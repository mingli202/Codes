from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class Report:
    report_date: date
    data: dict[str, float | None]  # product -> qty (None allowed)


def compile_json_reports_to_excel(
    *,
    json_root: str | Path = "json",
    out_path: str | Path = "compiled_sales.xlsx",
) -> Path:
    """
    Reads JSON reports from: json/<store_name>/<YYYYMMDD>.json
    Each JSON is a hashmap: { "<product>": <quantity> }

    Writes an Excel workbook:
      - 1 sheet per store
      - Column A: Product
      - Columns B..: report dates (sorted oldest -> newest), displayed as MM/DD
      - Header row above dates groups by year (merged cells)
      - Last column: row total across dates
      - Last row: column total across products (incl bottom-right grand total)
      - Missing products in a report => blank cell
    """
    json_root = Path(json_root)
    out_path = Path(out_path)

    stores = _load_all_stores(json_root)

    wb = Workbook()
    # Remove the default sheet
    default = wb.active
    wb.remove(default)

    for store_name in sorted(stores.keys(), key=str.lower):
        reports = stores[store_name]
        if not reports:
            continue
        _write_store_sheet(wb, store_name, reports)

    wb.save(out_path)
    return out_path


# ----------------------------
# Loading
# ----------------------------


def _load_all_stores(json_root: Path) -> dict[str, list[Report]]:
    if not json_root.exists() or not json_root.is_dir():
        raise ValueError(f"json_root must be an existing directory: {json_root}")

    stores: dict[str, list[Report]] = {}

    for store_dir in sorted([p for p in json_root.iterdir() if p.is_dir()]):
        reports: list[Report] = []
        for fp in sorted(store_dir.glob("*.json")):
            d = _parse_report_date(fp.stem)
            if d is None:
                continue

            raw = json.loads(fp.read_text(encoding="utf-8"))
            if not isinstance(raw, dict):
                raise ValueError(f"Expected dict JSON in {fp}, got {type(raw)}")

            cleaned: dict[str, float | None] = {}
            for k, v in raw.items():
                if k is None:
                    continue
                name = str(k).strip()
                if not name:
                    continue
                cleaned[name] = _to_number_or_none(v)

            reports.append(Report(report_date=d, data=cleaned))

        # sort oldest -> newest
        reports.sort(key=lambda r: r.report_date)
        stores[store_dir.name] = reports

    return stores


def _parse_report_date(stem: str) -> date | None:
    # Expect YYYYMMDD
    try:
        return datetime.strptime(stem, "%Y%m%d").date()
    except ValueError:
        return None


def _to_number_or_none(v: Any) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("\u00a0", " ").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


# ----------------------------
# Writing
# ----------------------------


def _safe_sheet_name(name: str) -> str:
    # Excel sheet constraints: max 31 chars, no : \ / ? * [ ]
    bad = set(r"[]:*?/\\")
    cleaned = "".join("_" if ch in bad else ch for ch in name).strip()
    if not cleaned:
        cleaned = "Sheet"
    return cleaned[:31]


def _write_store_sheet(wb: Workbook, store_name: str, reports: list[Report]) -> None:
    ws = wb.create_sheet(title=_safe_sheet_name(store_name))

    dates = [r.report_date for r in reports]

    # Union of products across all reports
    products_set = set()
    for r in reports:
        products_set.update(r.data.keys())
    products = sorted(products_set, key=str.lower)

    # Layout constants
    header_year_row = 1
    header_date_row = 2
    first_data_row = 3

    product_col = 1
    first_date_col = 2
    last_date_col = first_date_col + len(dates) - 1
    total_col = last_date_col + 1

    # Styles
    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill("solid", fgColor="F2F2F2")

    # Column headers
    ws.cell(row=header_date_row, column=product_col, value="Product").font = bold
    ws.cell(row=header_date_row, column=product_col).alignment = center
    ws.cell(row=header_date_row, column=product_col).fill = header_fill

    ws.cell(row=header_date_row, column=total_col, value="TOTAL").font = bold
    ws.cell(row=header_date_row, column=total_col).alignment = center
    ws.cell(row=header_date_row, column=total_col).fill = header_fill

    # Date headers (row 2), written as actual dates but formatted as MM/DD
    for i, d in enumerate(dates):
        c = first_date_col + i
        cell = ws.cell(row=header_date_row, column=c, value=d)
        cell.number_format = "mm/dd"
        cell.font = bold
        cell.alignment = center
        cell.fill = header_fill

    # Year group header (row 1) with merged cells over contiguous same-year blocks
    ws.cell(row=header_year_row, column=product_col, value="").fill = header_fill
    ws.cell(row=header_year_row, column=total_col, value="").fill = header_fill

    year_runs: list[tuple[int, int, int]] = []  # (year, start_col, end_col)
    run_year = dates[0].year
    run_start = first_date_col
    for i, d in enumerate(dates):
        col = first_date_col + i
        if d.year != run_year:
            year_runs.append((run_year, run_start, col - 1))
            run_year = d.year
            run_start = col
    year_runs.append((run_year, run_start, last_date_col))

    for y, c1, c2 in year_runs:
        ws.merge_cells(
            start_row=header_year_row,
            start_column=c1,
            end_row=header_year_row,
            end_column=c2,
        )
        cell = ws.cell(row=header_year_row, column=c1, value=str(y))
        cell.font = bold
        cell.alignment = center
        # apply fill to all cells in merged span
        for c in range(c1, c2 + 1):
            ws.cell(row=header_year_row, column=c).fill = header_fill

    # Freeze panes below headers and after product column
    ws.freeze_panes = ws.cell(row=first_data_row, column=first_date_col)

    # Data rows
    for r_idx, product in enumerate(products):
        excel_row = first_data_row + r_idx

        pc = ws.cell(row=excel_row, column=product_col, value=product)
        pc.alignment = left

        # Fill quantities
        for i, rep in enumerate(reports):
            excel_col = first_date_col + i
            qty = rep.data.get(product, None)
            ws.cell(
                row=excel_row, column=excel_col, value=qty if qty is not None else None
            )

        # Row total formula (sum across dates)
        start_letter = get_column_letter(first_date_col)
        end_letter = get_column_letter(last_date_col)
        ws.cell(
            row=excel_row,
            column=total_col,
            value=f"=SUM({start_letter}{excel_row}:{end_letter}{excel_row})",
        )

    # Totals row at bottom
    totals_row = first_data_row + len(products)
    ws.cell(row=totals_row, column=product_col, value="TOTAL").font = bold
    ws.cell(row=totals_row, column=product_col).alignment = left
    ws.cell(row=totals_row, column=product_col).fill = header_fill

    # Column totals per date
    for c in range(first_date_col, last_date_col + 1):
        col_letter = get_column_letter(c)
        ws.cell(
            row=totals_row,
            column=c,
            value=f"=SUM({col_letter}{first_data_row}:{col_letter}{totals_row - 1})",
        ).font = bold
        ws.cell(row=totals_row, column=c).fill = header_fill

    # Bottom-right grand total (sum of totals row across all date columns)
    start_letter = get_column_letter(first_date_col)
    end_letter = get_column_letter(last_date_col)
    ws.cell(
        row=totals_row,
        column=total_col,
        value=f"=SUM({start_letter}{totals_row}:{end_letter}{totals_row})",
    ).font = bold
    ws.cell(row=totals_row, column=total_col).fill = header_fill

    # Make header rows taller
    ws.row_dimensions[header_year_row].height = 18
    ws.row_dimensions[header_date_row].height = 18

    # Basic column sizing
    ws.column_dimensions[get_column_letter(product_col)].width = 45
    for c in range(first_date_col, total_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10

    # Center numeric area (dates + totals)
    for r in range(first_data_row, totals_row + 1):
        for c in range(first_date_col, total_col + 1):
            ws.cell(row=r, column=c).alignment = center

    # Apply fill to TOTAL header cell area (row1/row2 total col)
    ws.cell(row=header_year_row, column=total_col).fill = header_fill
    ws.cell(row=header_date_row, column=total_col).fill = header_fill


# ----------------------------
# Example usage
# ----------------------------
if __name__ == "__main__":
    compile_json_reports_to_excel(
        json_root="json",
        out_path="weekly_sales_compiled.xlsx",
    )
